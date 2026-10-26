"""
Full validation report (validation_report.build_report) -> Word document (.docx).

Same content as the PDF. Contents and list of tables are Word fields that fill in when the document is opened
(Word asks to update the fields). The trend chart is PDF-only; its data is in the table above it.
"""
import io
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x2F, 0x5D, 0x8A)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
RATING_COLORS = {"Satisfactory": RGBColor(0x1E, 0x7B, 0x4A), "Needs attention": RGBColor(0xB2, 0x6A, 0x00),
                 "Satisfactory with recommendations": RGBColor(0xB2, 0x6A, 0x00)}


# ==================================================================
# Low-level helpers
# ==================================================================

def _field(paragraph, instruction, placeholder=""):
    """Insert a Word field (TOC, SEQ, PAGE ...) into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _update_fields_on_open(document):
    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _runs(paragraph, text, size=None, color=None):
    """Add text with **bold** markup as runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bool(i % 2)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    tc_pr.append(shading)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _widths(header, rows, total_cm):
    """Column widths from the longest word / value (same idea as the PDF)."""
    cols = list(zip(*([header] + rows)))
    word = [max(len(w) for cell in col for w in (str(cell).split() or [""])) for col in cols]
    value = [max((len(str(c)) for c in col[1:]), default=0) for col in cols]
    wanted = [max(max(w, 4), min(v, 45)) for w, v in zip(word, value)]
    return [Cm(total_cm * w / sum(wanted)) for w in wanted]


# ==================================================================
# Document
# ==================================================================

def _styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
    caption = document.styles["Caption"]
    caption.font.size = Pt(9)
    caption.font.bold = True
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0x1F, 0x29, 0x33)


def _cover(document, meta):
    for _ in range(5):
        document.add_paragraph()
    p = document.add_paragraph()
    _runs(p, meta["organisation"].upper(), size=13, color=ACCENT)
    p.runs[0].bold = True
    p = document.add_paragraph()
    run = p.add_run(meta["title"])
    run.bold, run.font.size = True, Pt(26)
    p = document.add_paragraph()
    _runs(p, f"Prepared by {meta['validator']}", size=11, color=MUTED)
    document.add_paragraph()
    table = document.add_table(rows=0, cols=2)
    for label, value in (("Overall rating", meta.get("overall_rating") or "-"),
                         ("Latest period assessed", meta.get("period") or "-"),
                         ("Report date", meta["generated_at"]), ("AI model", meta.get("model") or "-"),
                         ("AI text verified", f"{meta.get('ai_parts', 0)} of {meta.get('parts', 0)} parts")):
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        run = cells[1].paragraphs[0].add_run(value)
        if label == "Overall rating":
            run.bold = True
            run.font.color.rgb = RATING_COLORS.get(value, MUTED)
    for _ in range(4):
        document.add_paragraph()
    p = document.add_paragraph()
    _runs(p, "This report was drafted by an AI agent from evidence produced by deterministic validation code. All "
             "tables, figures and ratings are computed by code, and every number in the AI text was checked against "
             "that evidence. The report is advisory until signed off by the validator.", size=9, color=MUTED)


def _footer(section, meta):
    p = section.footer.paragraphs[0]
    _runs(p, f"{meta['title']} - generated {meta['generated_at']} - page ", size=8, color=MUTED)
    _field(p, "PAGE", "1")
    h = section.header.paragraphs[0]
    _runs(h, meta["organisation"], size=8, color=MUTED)


def _table(document, block, number_holder):
    caption = document.add_paragraph(style="Caption")
    caption.add_run("Table ")
    _field(caption, "SEQ Table \\* ARABIC", str(number_holder[0]))
    caption.add_run(f": {block['caption']}")
    caption.paragraph_format.keep_with_next = True
    header, rows = block["header"], block["rows"]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = _widths(header, rows, 16.5)
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        _runs(cell.paragraphs[0], text, size=8.5)
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, "EEF1F5")
    _repeat_header(table.rows[0])
    for r in rows:
        cells = table.add_row().cells
        for i, text in enumerate(r):
            _runs(cells[i].paragraphs[0], text, size=8.5)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _signoff(document, approval):
    rows = [("Validator name", ""), ("Decision", "Approve  /  Request changes  /  Reject"), ("Comments", ""),
            ("Date", ""), ("Signature", "")]
    if approval:
        rows = [("Validator name", approval.get("reviewer") or "-"), ("Decision", approval.get("decision") or "-"),
                ("Comments", approval.get("comment") or "-"), ("Date", approval.get("run_at") or "-")]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(value)
        cells[0].width, cells[1].width = Cm(4), Cm(12.5)
        if not approval:
            cells[1].paragraphs[0].paragraph_format.space_after = Pt(18)


def render_full_docx(report, approval=None):
    """report: validation_report.build_report(...) result. Returns .docx bytes."""
    meta = report["meta"]
    document = Document()
    _styles(document)
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(2.0)
    _cover(document, meta)
    body = document.add_section(WD_SECTION.NEW_PAGE)
    body.footer.is_linked_to_previous = False
    body.header.is_linked_to_previous = False
    _footer(body, meta)
    _update_fields_on_open(document)

    numbers, table_no = [0, 0, 0], [0]
    for block in report["blocks"]:
        kind = block["type"]
        if kind == "front_matter":
            document.add_paragraph("Contents", style="Heading 1")
            _field(document.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u',
                   "Right-click and choose Update Field to show the contents.")
            document.add_page_break()
            document.add_paragraph("List of Tables", style="Heading 1")
            _field(document.add_paragraph(), 'TOC \\h \\z \\c "Table"',
                   "Right-click and choose Update Field to show the list of tables.")
            document.add_page_break()
        elif kind == "heading":
            level = block["level"]
            if level == 0:
                document.add_paragraph(block["text"], style="Heading 1")
                continue
            numbers[level - 1] += 1
            for i in range(level, 3):
                numbers[i] = 0
            label = ".".join(str(n) for n in numbers[:level])
            document.add_paragraph(f"{label}  {block['text']}", style=f"Heading {min(level, 3)}")
        elif kind == "para":
            _runs(document.add_paragraph(), block["text"])
        elif kind == "label":
            p = document.add_paragraph()
            p.add_run(block["text"]).bold = True
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
        elif kind == "bullets":
            for text in block["items"]:
                _runs(document.add_paragraph(style="List Bullet"), text)
        elif kind == "numbered":
            for i, text in enumerate(block["items"], 1):
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                _runs(p, f"{i}.\t{text}")
        elif kind == "table":
            if block["rows"]:
                table_no[0] += 1
                _table(document, block, table_no)
        elif kind == "rating":
            p = document.add_paragraph()
            run = p.add_run(block["text"])
            run.bold = True
            run.font.color.rgb = RATING_COLORS.get(block["text"], MUTED)
        elif kind == "chart":
            p = document.add_paragraph()
            _runs(p, f"Chart: {block['caption']} (shown in the PDF version; the data is in the table above).",
                  size=8.5, color=MUTED)
        elif kind == "signoff":
            _signoff(document, approval)
        elif kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
