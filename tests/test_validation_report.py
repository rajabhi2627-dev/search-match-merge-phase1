"""Full validation report: guardrails, fallback and rendering, with a stand-in for the LLM (no API calls)."""
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd
import pytest

import data_generator as G
import matching as M
import performance as P
import report_docx as RD
import report_pdf as RP
import validation_report as R


@pytest.fixture(scope="module")
def data():
    records = G.generate_master_records()
    master = M.prepare_master(records)
    results = P.run_monitoring(master)
    production = {"Algorithm": "Production algorithm", **P.document_metrics(results)}
    benchmark = pd.DataFrame([production, {**production, "Algorithm": "Challenger (name)"}])
    return results, benchmark, pd.DataFrame(records), master


def _answer(group, evidence):
    """What a well-behaved model would submit: correct ratings and only numbers from the evidence."""
    rates = evidence.get("rule_based_ratings", {})
    if group in ("quant_a", "quant_b"):
        if group == "quant_a":
            first = evidence["performance_by_period"][0]
            point = f"The hit rate was {first['hit_rate_pct']}% in {first['period']}."
        else:
            last = evidence["errors_by_period"][-1]
            point = f"The Type II error rate was {last['type2_error_pct']}% in {last['period']}."
        return {part: {"analysis": [point, "The results are consistent across periods.", "No anomaly was found."],
                       "observation": "The review indicates results in line with the rule-based assessment.",
                       "recommendation": "The model owner may consider continued monitoring.", "rating": rating}
                for part, rating in rates.items()}
    if group == "qual_a":
        def rows(areas):
            return [{"area": a, "observation": "Documented in the README.", "recommendation": "Keep it current."}
                    for a in areas]
        return {"documentation_review": rows(R.DOC_AREAS), "data_quality_rows": rows(R.DQ_AREAS),
                "data_quality_conclusion": "Data quality is adequate."}
    if group == "qual_b":
        return {k: {"assessment": "Reviewed.", "observation": "Documented.", "recommendation": "Keep it current."}
                for k in R.QUAL_SECTIONS}
    return {"qualitative_assessment": "Sound.", "quantitative_assessment": "Mixed.", "conclusion": "See ratings.",
            "validation_summary": ["Work performed.", "Results."], "overall_rating": evidence["overall_rating"]}


def good_ask(group, evidence, check):
    answer = _answer(group, evidence)
    return answer, "test-model", 0, check(answer)


def bad_ask(group, evidence, check):
    answer = json.loads(json.dumps(_answer(group, evidence)))
    if group == "quant_a":
        answer["performance"]["observation"] = "The hit rate was 77.7%."                 # invented number
        part = "challenger"
        answer[part]["rating"] = R.NEEDS if answer[part]["rating"] == R.SAT else R.SAT   # wrong rating
    if group == "qual_a":
        answer["documentation_review"] = answer["documentation_review"][:2]              # missing areas
    if group == "qual_b":
        answer["governance"]["observation"] = "See Table 3."                             # table reference
    return answer, "test-model", 1, check(answer)


def test_correct_answers_pass_every_check(data):
    report = R.build_report(*data, ask=good_ask)
    assert report["meta"]["ai_parts"] == report["meta"]["parts"] == len(report["log"])


def test_mistakes_are_caught_and_replaced(data):
    report = R.build_report(*data, ask=bad_ask)
    failed = {row["part"].split(" - ")[-1] for row in report["log"] if row["result"] != "AI text verified"}
    assert failed == {"performance", "challenger", "documentation review", "governance"}
    body = [b for b in report["blocks"] if b.get("caption") != "AI generation and verification log"]
    text = json.dumps(body, default=str)  # the log in Annexure B quotes the rejected text on purpose
    assert "77.7%" not in text and "See Table 3" not in text


def test_ratings_come_from_rules(data):
    ev = R.collect(*data)
    rates = R.ratings(ev)
    assert set(rates.values()) <= {R.SAT, R.NEEDS}
    assert R.overall_rating(ev, rates) in R.OVERALL_RATINGS


def test_report_renders_as_pdf_and_word(data):
    report = R.build_report(*data, ask=good_ask)
    approval = {"reviewer": "A. Validator", "decision": "Approve", "comment": "ok", "run_at": "2026-10-26T12:00"}
    assert RP.render_full_pdf(report, approval).startswith(b"%PDF")
    import docx
    document = docx.Document(io.BytesIO(RD.render_full_docx(report, approval)))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "1  Background" in headings and "6  Validation Summary" in headings
    assert len(document.tables) > 20
